from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.responses import FileResponse
from fastapi.middleware.cors import CORSMiddleware
from pathlib import Path
from PIL import Image, ImageOps
import fitz, pypdf, subprocess, tempfile, shutil, os, uuid, zipfile
from pdf2docx import Converter
from pptx import Presentation

app = FastAPI(title='USR Advisory Free Tools API', version='1.0.0')
app.add_middleware(CORSMiddleware, allow_origins=['*'], allow_methods=['POST','GET'], allow_headers=['*'])
ROOT = Path(tempfile.gettempdir()) / 'usr-tools-api'
ROOT.mkdir(exist_ok=True)
MAX_FILE = 100 * 1024 * 1024

def jobdir():
    p = ROOT / uuid.uuid4().hex
    p.mkdir()
    return p

def safe_name(n):
    return ''.join(c for c in (n or 'file') if c.isalnum() or c in '._-') or 'file'

async def save_upload(upload, folder):
    data = await upload.read()
    if len(data) > MAX_FILE: raise HTTPException(413, 'File is too large. Maximum size is 100 MB per file.')
    p = folder / safe_name(upload.filename)
    p.write_bytes(data)
    return p

def soffice(src, outdir, fmt):
    r = subprocess.run(['soffice','--headless','--convert-to',fmt,'--outdir',str(outdir),str(src)], capture_output=True, text=True, timeout=180)
    if r.returncode != 0: raise RuntimeError(r.stderr or r.stdout or 'LibreOffice conversion failed')
    expected = outdir / (src.stem + '.' + fmt.split(':')[0])
    if expected.exists(): return expected
    candidates = list(outdir.glob(src.stem + '.*'))
    if not candidates: raise RuntimeError('Conversion produced no output')
    return candidates[0]

def response(path, media=None):
    return FileResponse(path, media_type=media, filename=path.name)

@app.get('/api/health')
def health(): return {'ok': True, 'service': 'USR Free Tools API'}

@app.post('/api/word-to-pdf')
async def word_to_pdf(file: UploadFile = File(...)):
    d=jobdir(); src=await save_upload(file,d); out=soffice(src,d,'pdf'); return response(out,'application/pdf')

@app.post('/api/pdf-to-word')
async def pdf_to_word(file: UploadFile = File(...)):
    d=jobdir(); src=await save_upload(file,d); out=d/'converted.docx'
    cv=Converter(str(src)); cv.convert(str(out), start=0, end=None); cv.close(); return response(out,'application/vnd.openxmlformats-officedocument.wordprocessingml.document')

@app.post('/api/ppt-to-word')
async def ppt_to_word(file: UploadFile = File(...)):
    d=jobdir(); src=await save_upload(file,d); pdf=soffice(src,d,'pdf'); out=d/'converted.docx'
    cv=Converter(str(pdf)); cv.convert(str(out), start=0, end=None); cv.close(); return response(out,'application/vnd.openxmlformats-officedocument.wordprocessingml.document')

@app.post('/api/word-to-ppt')
async def word_to_ppt(file: UploadFile = File(...)):
    d=jobdir(); src=await save_upload(file,d); pdf=soffice(src,d,'pdf'); doc=fitz.open(pdf); prs=Presentation(); prs.slide_width=13.333*914400; prs.slide_height=7.5*914400
    blank=prs.slide_layouts[6]
    for page in doc:
        pix=page.get_pixmap(matrix=fitz.Matrix(2.2,2.2), alpha=False)
        img=d/f'page-{page.number+1}.png'; pix.save(img)
        slide=prs.slides.add_slide(blank); slide.shapes.add_picture(str(img),0,0,width=prs.slide_width,height=prs.slide_height)
    out=d/'word-to-ppt.pptx'; prs.save(out); return response(out,'application/vnd.openxmlformats-officedocument.presentationml.presentation')

@app.post('/api/merge-pdf')
async def merge_pdf(files: list[UploadFile] = File(...)):
    if not files: raise HTTPException(400,'Select at least two PDF files.')
    d=jobdir(); writer=pypdf.PdfWriter()
    for f in files:
        src=await save_upload(f,d); reader=pypdf.PdfReader(str(src));
        for page in reader.pages: writer.add_page(page)
    out=d/'merged.pdf';
    with open(out,'wb') as h: writer.write(h)
    return response(out,'application/pdf')

@app.post('/api/compress-pdf')
async def compress_pdf(file: UploadFile = File(...), compression: str = Form('recommended')):
    d=jobdir(); src=await save_upload(file,d); out=d/'compressed.pdf'
    preset={'recommended':'/ebook','high':'/screen','extreme':'/screen'}.get(compression,'/ebook')
    r=subprocess.run(['gs','-sDEVICE=pdfwrite','-dCompatibilityLevel=1.7','-dNOPAUSE','-dBATCH','-dSAFER',f'-dPDFSETTINGS={preset}','-sOutputFile='+str(out),str(src)],capture_output=True,text=True,timeout=240)
    if r.returncode!=0: raise RuntimeError(r.stderr or 'PDF compression failed')
    return response(out,'application/pdf')

@app.post('/api/pdf-to-jpg')
async def pdf_to_jpg(file: UploadFile = File(...), dpi: int = Form(180)):
    return await pdf_images(file, dpi, 'jpg')

@app.post('/api/pdf-to-png')
async def pdf_to_png(file: UploadFile = File(...), dpi: int = Form(180)):
    return await pdf_images(file, dpi, 'png')

async def pdf_images(file,dpi,fmt):
    d=jobdir(); src=await save_upload(file,d); doc=fitz.open(src); entries=[]; scale=max(72,min(300,int(dpi)))/72
    for i,p in enumerate(doc):
        pix=p.get_pixmap(matrix=fitz.Matrix(scale,scale), alpha=False)
        ext='jpg' if fmt=='jpg' else 'png'; out=d/f'page-{i+1}.{ext}'; pix.save(out,output=ext); entries.append(out)
    if len(entries)==1: return response(entries[0], 'image/jpeg' if fmt=='jpg' else 'image/png')
    z=d/f'pdf-pages-{fmt}.zip'
    with zipfile.ZipFile(z,'w',zipfile.ZIP_DEFLATED) as zz:
        for p in entries: zz.write(p,p.name)
    return response(z,'application/zip')

@app.post('/api/jpg-to-pdf')
async def jpg_to_pdf(files: list[UploadFile] = File(...), page_size: str = Form('A4'), orientation: str = Form('auto'), margin: int = Form(10)):
    return await images_pdf(files,page_size,orientation,margin)

@app.post('/api/images-to-pdf')
async def images_to_pdf(files: list[UploadFile] = File(...), page_size: str = Form('A4'), orientation: str = Form('auto'), margin: int = Form(10)):
    return await images_pdf(files,page_size,orientation,margin)

async def images_pdf(files,page_size,orientation,margin):
    d=jobdir(); out=d/'images.pdf'; imgs=[]
    for f in files: imgs.append(await save_upload(f,d))
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import A4, LETTER, A5
    from reportlab.lib.utils import ImageReader
    sizes={'A4':A4,'LETTER':LETTER,'A5':A5}; base=sizes.get(page_size.upper(),A4)
    c=None
    for idx,p in enumerate(imgs):
        with Image.open(p) as im: w,h=im.size
        land=orientation.lower()=='landscape' or (orientation.lower()=='auto' and w>=h)
        pw,ph=(base[1],base[0]) if land else base
        if c is None: c=canvas.Canvas(str(out),pagesize=(pw,ph))
        else: c.setPageSize((pw,ph))
        maxw=pw-2*margin*2.83465; maxh=ph-2*margin*2.83465
        iw=w; ih=h; s=min(maxw/iw,maxh/ih); dw=iw*s; dh=ih*s
        x=(pw-dw)/2; y=(ph-dh)/2
        c.drawImage(ImageReader(str(p)),x,y,width=dw,height=dh,preserveAspectRatio=True,mask='auto')
        c.showPage()
    if c is None: raise HTTPException(400,'Select at least one image.')
    c.save(); return response(out,'application/pdf')

@app.post('/api/image-compressor')
async def image_compressor(files:list[UploadFile]=File(...), target_kb:int=Form(200), max_width:int=Form(2400), fmt:str=Form('jpg')):
    d=jobdir(); outs=[]
    for f in files:
        src=await save_upload(f,d)
        with Image.open(src) as im:
            im=ImageOps.exif_transpose(im).convert('RGB');
            if im.width>max_width: im.thumbnail((max_width,max_width),Image.Resampling.LANCZOS)
            ext='webp' if fmt=='webp' else 'jpg'; out=d/(src.stem+'.'+ext)
            q=90
            while q>=25:
                im.save(out,'WEBP' if ext=='webp' else 'JPEG',quality=q,optimize=True,progressive=True)
                if out.stat().st_size<=target_kb*1024: break
                q-=5
        outs.append(out)
    if len(outs)==1:return response(outs[0],'image/webp' if fmt=='webp' else 'image/jpeg')
    z=d/'compressed-images.zip';
    with zipfile.ZipFile(z,'w',zipfile.ZIP_DEFLATED) as zz:
        for p in outs: zz.write(p,p.name)
    return response(z,'application/zip')

@app.post('/api/image-resizer')
async def image_resizer(files:list[UploadFile]=File(...), width:int=Form(...), height:int=Form(...), keep:bool=Form(True)):
    d=jobdir(); outs=[]
    for f in files:
        src=await save_upload(f,d)
        with Image.open(src) as im:
            im=ImageOps.exif_transpose(im)
            nw,nh=width,height
            if keep:
                s=min(width/im.width,height/im.height); nw=max(1,round(im.width*s)); nh=max(1,round(im.height*s))
            out=d/(src.stem+'-resized.jpg'); im.convert('RGB').resize((nw,nh),Image.Resampling.LANCZOS).save(out,'JPEG',quality=95,optimize=True)
        outs.append(out)
    if len(outs)==1:return response(outs[0],'image/jpeg')
    z=d/'resized-images.zip';
    with zipfile.ZipFile(z,'w',zipfile.ZIP_DEFLATED) as zz:
        for p in outs: zz.write(p,p.name)
    return response(z,'application/zip')
